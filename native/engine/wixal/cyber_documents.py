"""Offline document derivatives of the already redacted report snapshot."""
from io import BytesIO
from xml.sax.saxutils import escape


def blocks(report, executive=False):
    from .cyber_reports import markdown
    code=[];inside=False
    for line in markdown(report,executive).splitlines():
        if line.startswith('```'):
            if inside:yield 'code','\n'.join(code);code=[]
            inside=not inside;continue
        if inside:code.append(line)
        elif line.strip():yield 'text',line


def pdf(report, executive=False):
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Preformatted
    stream=BytesIO();styles=getSampleStyleSheet()
    for style in styles.byName.values():style.textColor=colors.black
    styles['BodyText'].fontSize=10;styles['BodyText'].leading=14
    styles['BodyText'].wordWrap='CJK';styles['BodyText'].alignment=TA_LEFT
    story=[]
    styles['Code'].fontSize=8;styles['Code'].leading=10
    for kind,line in blocks(report,executive):
        if kind=='code':
            story.extend((Preformatted(line,styles['Code'],maxLineLength=100),Spacer(1,6)));continue
        level=2 if line.startswith('## ') else 1 if line.startswith('# ') else 0
        text=line[level+1:] if level else line
        story.extend((Paragraph(escape(text),styles['Heading'+str(level)] if level else styles['BodyText']),Spacer(1,6)))
    def footer(canvas,document):
        canvas.setFont('Helvetica',8);canvas.drawString(42,25,'Wixal assessment - review before sharing')
        canvas.drawRightString(A4[0]-42,25,str(document.page))
    SimpleDocTemplate(stream,pagesize=A4,rightMargin=42,leftMargin=42,topMargin=42,bottomMargin=42,
                      title=report['engagement']['name'],author='Wixal').build(story,onFirstPage=footer,onLaterPages=footer)
    return stream.getvalue()


def docx(report, executive=False):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    document=Document();section=document.sections[0]
    section.top_margin=section.bottom_margin=Inches(.7)
    section.left_margin=section.right_margin=Inches(.7)
    for name in ('Normal','Title','Heading 1','Heading 2'):
        style=document.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
    for element in document.styles.element.xpath('.//w:pBdr'):
        element.getparent().remove(element)
    document.styles['Normal'].font.size=Pt(10)
    for kind,line in blocks(report,executive):
        if kind=='code':
            paragraph=document.add_paragraph();paragraph.paragraph_format.space_after=Pt(6)
            paragraph.paragraph_format.line_spacing=1
            run=paragraph.add_run(line);run.font.name='Courier New';run.font.size=Pt(8)
            continue
        if line.startswith('# '):document.add_paragraph(line[2:],'Title')
        elif line.startswith('## '):document.add_heading(line[3:],1)
        else:document.add_paragraph(line)
    section.footer.paragraphs[0].text='Wixal assessment - review before sharing'
    document.core_properties.title=report['engagement']['name'];document.core_properties.author='Wixal'
    stream=BytesIO();document.save(stream);return stream.getvalue()
